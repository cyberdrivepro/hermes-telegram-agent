"""Local, bounded artifact builders for the Omega capability registry.

These functions create reviewable files; they do not publish forms, fetch media,
install packages, transcribe speech, or certify mechanical load capacity.
Optional dependencies are loaded only by the capability that needs them.

API references (checked 2026-09-07):
https://developers.google.com/workspace/forms/api/reference/rest/v1/forms
https://developers.google.com/apps-script/reference/forms/scale-item
https://developers.google.com/apps-script/reference/forms/multiple-choice-item
https://openpyxl.readthedocs.io/en/stable/simple_formulae.html
https://ffmpeg.org/ffmpeg.html
https://ffmpeg.org/ffmpeg-filters.html#subtitles-1
https://ffmpeg.org/ffmpeg-protocols.html
"""

from __future__ import annotations

import html
import json
import math
import re
import shutil
import struct
import subprocess
import tempfile
import uuid
from collections import Counter
from pathlib import Path
from typing import Any, Callable


class ArtifactDependencyError(RuntimeError):
    """A deployment dependency required for a requested artifact is absent."""


def _text(value: Any, name: str, maximum: int = 2000, empty: bool = False) -> str:
    if not isinstance(value, str) or len(value) > maximum or (not empty and not value.strip()):
        raise ValueError(f"{name} must be {'a' if empty else 'a nonempty'} string of at most {maximum} characters")
    if any(ord(char) < 32 and char not in "\n\t" for char in value):
        raise ValueError(f"{name} contains unsupported control characters")
    return value


def _number(value: Any, name: str, minimum: float, maximum: float) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{name} must be a number")
    try:
        number = float(value)
    except OverflowError as exc:
        raise ValueError(f"{name} is outside the supported range") from exc
    if not math.isfinite(number) or not minimum <= number <= maximum:
        raise ValueError(f"{name} must be between {minimum} and {maximum}")
    return number


def _integer(value: Any, name: str, minimum: int, maximum: int) -> int:
    number = _number(value, name, minimum, maximum)
    if not number.is_integer():
        raise ValueError(f"{name} must be an integer")
    return int(number)


def _boolean(value: Any, name: str) -> bool:
    if not isinstance(value, bool):
        raise ValueError(f"{name} must be a boolean")
    return value


def _folder(workspace: Path, prefix: str) -> Path:
    workspace = Path(workspace).resolve()
    if not workspace.is_dir():
        raise ValueError("workspace must be an existing directory")
    directory = workspace / f"{prefix}_{uuid.uuid4().hex[:12]}"
    directory.mkdir()
    return directory


def _result(summary: str, data: dict, directory: Path, workspace: Path) -> dict:
    return {"summary": summary, "data": data, "artifacts": [
        path.relative_to(Path(workspace).resolve()).as_posix()
        for path in sorted(directory.iterdir()) if path.is_file()
    ]}


def _json_write(path: Path, data: Any) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


_QUESTION_TYPES = {"short_text", "paragraph", "multiple_choice", "checkboxes", "dropdown", "scale"}


def _survey_spec(payload: dict) -> tuple[dict, list[dict], int]:
    title = _text(payload.get("title", "Employee feedback"), "title", 200)
    description = _text(payload.get("description", ""), "description", 3000, empty=True)
    questions = payload.get("questions", [
        {"id": "satisfaction", "title": "Overall work satisfaction", "type": "scale", "required": True},
        {"id": "support", "title": "Do you have the support you need?", "type": "multiple_choice", "options": ["Yes", "Sometimes", "No"]},
        {"id": "improvements", "title": "What would improve your experience?", "type": "paragraph"},
    ])
    if not isinstance(questions, list) or not 1 <= len(questions) <= 50:
        raise ValueError("questions must contain between 1 and 50 objects")
    normalized, identifiers = [], set()
    for i, question in enumerate(questions):
        if not isinstance(question, dict):
            raise ValueError("Every question must be an object")
        identifier = question.get("id", f"q{i + 1}")
        if not isinstance(identifier, str) or not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,39}", identifier):
            raise ValueError("Question ids must start with a letter and contain at most 40 letters, digits, or underscores")
        if identifier in identifiers:
            raise ValueError(f"Duplicate question id: {identifier}")
        identifiers.add(identifier)
        kind = question.get("type", "short_text")
        if not isinstance(kind, str) or kind not in _QUESTION_TYPES:
            raise ValueError(f"Supported question types: {', '.join(sorted(_QUESTION_TYPES))}")
        item = {"id": identifier, "title": _text(question.get("title"), "question title", 500),
                "type": kind, "required": _boolean(question.get("required", False), "required")}
        if kind in {"multiple_choice", "checkboxes", "dropdown"}:
            options = question.get("options")
            if not isinstance(options, list) or not 1 <= len(options) <= 50:
                raise ValueError("Choice questions need 1 to 50 options")
            item["options"] = [_text(option, "option", 200) for option in options]
            if len(set(item["options"])) != len(options):
                raise ValueError("Choice options must be unique")
        if kind == "scale":
            item["min"] = _integer(question.get("min", 1), "scale min", 0, 1)
            item["max"] = _integer(question.get("max", 5), "scale max", 3, 10)
        normalized.append(item)
    responses = payload.get("responses", [])
    if not isinstance(responses, list) or len(responses) > 2000:
        raise ValueError("responses must be a list of at most 2000 objects")
    cleaned = []
    for response in responses:
        if not isinstance(response, dict) or set(response) - identifiers:
            raise ValueError("Each response must be an object keyed by known question ids")
        row = {}
        for item in normalized:
            value = response.get(item["id"], None)
            missing = value is None or value == "" or value == []
            if missing:
                if item["required"]:
                    raise ValueError(f"Missing required response: {item['id']}")
                row[item["id"]] = None
                continue
            if item["type"] == "scale":
                value = _integer(value, item["id"], item["min"], item["max"])
            elif item["type"] == "checkboxes":
                if not isinstance(value, list) or len(value) > 50 or any(not isinstance(v, str) or v not in item["options"] for v in value):
                    raise ValueError(f"Invalid checkbox choices for {item['id']}")
                if len(set(value)) != len(value):
                    raise ValueError("Checkbox answers must not repeat a choice")
            else:
                value = _text(value, item["id"], 2048)
                if "options" in item and value not in item["options"]:
                    raise ValueError(f"Unknown choice for {item['id']}")
            row[item["id"]] = value
        cleaned.append(row)
    capacity = _integer(payload.get("analysis_capacity", max(500, len(cleaned))), "analysis_capacity", max(1, len(cleaned)), 5000)
    return {"title": title, "description": description, "questions": normalized}, cleaned, capacity


def _form_requests(spec: dict) -> dict:
    requests = [{"updateFormInfo": {"info": {"description": spec["description"]}, "updateMask": "description"}}]
    for i, item in enumerate(spec["questions"]):
        question: dict = {"required": item["required"]}
        if item["type"] in {"short_text", "paragraph"}:
            question["textQuestion"] = {"paragraph": item["type"] == "paragraph"}
        elif item["type"] == "scale":
            question["scaleQuestion"] = {"low": item["min"], "high": item["max"]}
        else:
            question["choiceQuestion"] = {
                "type": {"multiple_choice": "RADIO", "checkboxes": "CHECKBOX", "dropdown": "DROP_DOWN"}[item["type"]],
                "options": [{"value": value} for value in item["options"]], "shuffle": False,
            }
        requests.append({"createItem": {"item": {"title": item["title"], "questionItem": {"question": question}}, "location": {"index": i}}})
    return {"requests": requests}


def _apps_script(spec: dict) -> str:
    # ensure_ascii also escapes Unicode line separators in JavaScript literals.
    return """// Generated draft. Running createSurvey creates a Form and response Sheet in your account.
// Review Google account permissions and publication settings before sharing.
function createSurvey() {
  const spec = """ + json.dumps(spec, ensure_ascii=True) + """;
  const form = FormApp.create(spec.title).setDescription(spec.description);
  spec.questions.forEach(q => {
    let item;
    if (q.type === 'short_text') item = form.addTextItem();
    else if (q.type === 'paragraph') item = form.addParagraphTextItem();
    else if (q.type === 'scale') item = form.addScaleItem().setBounds(q.min, q.max);
    else {
      if (q.type === 'multiple_choice') item = form.addMultipleChoiceItem();
      else if (q.type === 'checkboxes') item = form.addCheckboxItem();
      else item = form.addListItem();
      item.setChoiceValues(q.options);
    }
    item.setTitle(q.title).setRequired(q.required);
  });
  const sheet = SpreadsheetApp.create(spec.title + ' - Responses');
  form.setDestination(FormApp.DestinationType.SPREADSHEET, sheet.getId());
  const result = {editUrl: form.getEditUrl(), responseUrl: form.getPublishedUrl(), sheetUrl: sheet.getUrl()};
  console.log(JSON.stringify(result));
  return result;
}
"""


def _survey_html(spec: dict) -> str:
    blocks = []
    for item in spec["questions"]:
        name = item["id"]
        required = " required" if item["required"] else ""
        legend = html.escape(item["title"]) + (" *" if item["required"] else "")
        if item["type"] == "paragraph":
            control = f'<textarea name="{name}" maxlength="2048" aria-label="{html.escape(item["title"], quote=True)}"{required}></textarea>'
        elif item["type"] == "short_text":
            control = f'<input name="{name}" maxlength="2048" aria-label="{html.escape(item["title"], quote=True)}"{required}>'
        elif item["type"] in {"dropdown", "scale"}:
            options = item.get("options", list(range(item.get("min", 1), item.get("max", 5) + 1)))
            control = f'<select name="{name}" aria-label="{html.escape(item["title"], quote=True)}"{required}><option value="">Choose an answer</option>'
            control += "".join(f'<option value="{html.escape(str(value), quote=True)}">{html.escape(str(value))}</option>' for value in options) + "</select>"
        else:
            kind = "checkbox" if item["type"] == "checkboxes" else "radio"
            control = "".join(f'<label><input type="{kind}" name="{name}" value="{html.escape(value, quote=True)}"{required if kind == "radio" else ""}> {html.escape(value)}</label>' for value in item["options"])
        blocks.append(f'<fieldset data-id="{name}"><legend>{legend}</legend>{control}</fieldset>')
    # JSON inside script must escape < so user text cannot terminate the element.
    safe_json = json.dumps(spec, ensure_ascii=True).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    return """<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>""" + html.escape(spec["title"]) + """</title><style>
body{font:16px system-ui,sans-serif;background:#eef2f7;color:#163047;margin:0}main{max-width:760px;margin:32px auto;background:white;padding:32px;border-radius:18px}
h1{font-size:2rem}p{line-height:1.6}fieldset{border:1px solid #bdcddd;border-radius:10px;margin:20px 0;padding:18px}legend{font-weight:650;padding:0 8px}label{display:block;margin:12px 0}
input:not([type=radio]):not([type=checkbox]),textarea,select{box-sizing:border-box;width:100%;padding:12px;border:1px solid #91a6b8;border-radius:6px;font:inherit}textarea{min-height:110px}
button{background:#165a72;color:white;border:0;border-radius:8px;padding:14px 20px;font:inherit;cursor:pointer}progress{width:100%;height:14px}#status{min-height:24px}@media(max-width:600px){main{margin:8px;padding:18px}}
</style><main><h1>""" + html.escape(spec["title"]) + "</h1><p>" + html.escape(spec["description"]) + """</p>
<p>Your response stays in this browser. Export downloads a CSV; it sends nothing to a server.</p>
<progress id="progress" value="0" max="100" aria-label="Completion"></progress><form id="survey">""" + "".join(blocks) + """<button type="submit">Export response as CSV</button></form><p id="status" role="status"></p></main>
<script>
'use strict';
const spec = """ + safe_json + r""";
const form = document.getElementById('survey');
function answers() { const data = new FormData(form); return spec.questions.map(q => data.getAll(q.id).filter(v => v !== '')); }
function update() {
  const all = answers();
  spec.questions.forEach((q, i) => { if (q.type === 'checkboxes' && q.required) {
    form.querySelector('input[name="' + q.id + '"]').setCustomValidity(all[i].length ? '' : 'Choose at least one option.');
  }});
  document.getElementById('progress').value = 100 * all.filter(a => a.length).length / all.length;
}
function csvCell(value) {
  let text = String(value);
  if (/^[\s]*[=+\-@]/.test(text) || /^[\t\r]/.test(text)) text = "'" + text;
  return '"' + text.replace(/"/g, '""') + '"';
}
form.addEventListener('input', update);
form.addEventListener('submit', event => {
  event.preventDefault(); update(); if (!form.reportValidity()) return;
  const rows = [spec.questions.map(q => q.id), answers().map(a => a.join('; '))];
  const csv = '\uFEFF' + rows.map(row => row.map(csvCell).join(',')).join('\r\n');
  const url = URL.createObjectURL(new Blob([csv], {type:'text/csv;charset=utf-8'}));
  const link = document.createElement('a'); link.href = url; link.download = 'survey-response.csv'; link.click();
  setTimeout(() => URL.revokeObjectURL(url), 1000);
  document.getElementById('status').textContent = 'Your response CSV has been downloaded.';
});
update();
</script></html>"""


def _survey_workbook(spec: dict, responses: list[dict], capacity: int, path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.workbook.properties import CalcProperties
    except ImportError as exc:
        raise ArtifactDependencyError("survey.bundle with include_excel=true requires openpyxl in the deployment; include_excel=false creates the other survey files") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Responses"
    analysis = workbook.create_sheet("Analysis")
    guide = workbook.create_sheet("Guide")

    def literal(target, row, col, value):
        cell = target.cell(row, col, value)
        if isinstance(value, str):
            cell.data_type = "s"  # User answers, titles, and choices are never executable formulas.
        return cell

    literal(sheet, 1, 1, "Response ID")
    for i, item in enumerate(spec["questions"], 2):
        literal(sheet, 1, i, item["title"])
    for row_index, response in enumerate(responses, 2):
        sheet.cell(row_index, 1, row_index - 1)
        for col, item in enumerate(spec["questions"], 2):
            value = response[item["id"]]
            if isinstance(value, list):
                value = "; ".join(value)
            literal(sheet, row_index, col, value)
    analysis.append(["Survey analytics", "Value"])
    analysis.append(["Response count", f"=COUNT(Responses!A2:A{capacity + 1})"])
    analysis.append(["Question", "Answered", "Average rating"])
    for index, item in enumerate(spec["questions"], 4):
        col = get_column_letter(index - 2)
        span = f"Responses!{col}2:{col}{capacity + 1}"
        literal(analysis, index, 1, item["title"])
        analysis.cell(index, 2, f'=COUNTIF({span},"<>")')
        if item["type"] == "scale":
            analysis.cell(index, 3, f'=IF(COUNT({span})=0,"",AVERAGE({span}))').number_format = "0.00"
            validation = DataValidation(type="whole", operator="between", formula1=item["min"], formula2=item["max"], allow_blank=not item["required"])
            validation.showErrorMessage = True
            validation.errorTitle = "Rating outside the scale"
            validation.error = f"Enter a whole number from {item['min']} to {item['max']}."
            sheet.add_data_validation(validation)
            validation.add(f"{col}2:{col}{capacity + 1}")
    chart = BarChart()
    chart.title = "Responses by question"
    chart.y_axis.title = "Answered"
    chart.add_data(Reference(analysis, min_col=2, min_row=3, max_row=3 + len(spec["questions"])), titles_from_data=True)
    chart.set_categories(Reference(analysis, min_col=1, min_row=4, max_row=3 + len(spec["questions"])))
    chart.height, chart.width = 10, 20
    analysis.add_chart(chart, "E3")
    guide.append(["How to use this workbook", "Details"])
    guide.append(["Survey", spec["title"]])
    guide.append(["Entry", "Add one response per row in Responses; use a numeric Response ID for every row."])
    guide.append(["Capacity", f"Analysis formulas cover rows 2 through {capacity + 1}. Extend their ranges for additional data."])
    guide.append(["Recalculation", "Open in Excel or LibreOffice to calculate formulas. This generator does not evaluate Excel formulas or cache their results."])
    guide.append(["Google Forms", "The Apps Script creates a separate Google response Sheet; this XLSX is an offline analysis template, not a live connection."])
    guide.append(["Response columns", ", ".join(item["id"] for item in spec["questions"])])
    for row in guide:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for target in workbook:
        target.freeze_panes = "B2"
        target.sheet_view.showGridLines = False
        for cell in target[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="173F56")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        target.row_dimensions[1].height = 32
        for row in target.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EDF4F8")
        for col in range(1, target.max_column + 1):
            target.column_dimensions[get_column_letter(col)].width = 30 if col > 1 else 38
    guide.column_dimensions["B"].width = 100
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(spec['questions']) + 1)}{max(2, len(responses) + 1)}"
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    workbook.save(path)


def survey_bundle(payload: dict, workspace: Path) -> dict:
    """Build local Form API bodies, Apps Script, HTML, and an XLSX analysis template."""
    spec, responses, capacity = _survey_spec(payload)
    include_excel = _boolean(payload.get("include_excel", True), "include_excel")
    directory = _folder(workspace, "survey")
    try:
        if include_excel:
            _survey_workbook(spec, responses, capacity, directory / "analysis.xlsx")
        _json_write(directory / "survey.json", spec)
        _json_write(directory / "forms_create.json", {"info": {"title": spec["title"]}})
        _json_write(directory / "forms_batch_update.json", _form_requests(spec))
        (directory / "create_survey.gs").write_text(_apps_script(spec), encoding="utf-8")
        (directory / "survey.html").write_text(_survey_html(spec), encoding="utf-8")
        (directory / "README.txt").write_text(
            "Open survey.html locally to fill the survey and export one response as CSV.\n"
            "To create a Google Form, paste create_survey.gs into Apps Script and run createSurvey in your account.\n"
            "Alternatively, send forms_create.json to forms.create, then forms_batch_update.json to forms.batchUpdate using the returned formId and your credentials.\n"
            "No Google resource has been created or published by this local build.\n"
            "analysis.xlsx (when included) recalculates in Excel/LibreOffice and covers the stated capacity.\n", encoding="utf-8")
    except Exception:
        shutil.rmtree(directory)
        raise
    return _result(f"Created a {len(spec['questions'])}-question survey bundle.", {
        "question_count": len(spec["questions"]), "response_count": len(responses),
        "analysis_capacity": capacity, "excel_included": include_excel,
        "google_form_created": False, "formula_evaluation": "recalculates when opened in Excel/LibreOffice" if include_excel else None,
    }, directory, workspace)


def _plate_triangles(width: float, depth: float, thickness: float, spacing_x: float, spacing_y: float,
                     diameter: float, segments: int) -> list[tuple]:
    """Extrude a shared planar tessellation; boundary edges become side walls.

    Four mirrored rectangular cells each contain one circular hole. Rays through
    circle vertices and rectangle corners tessellate their annuli. Mirrored cells
    share identical edge vertices, preventing seams and T-junctions.
    """
    planar = []
    radius = diameter / 2
    for sx in (-1, 1):
        for sy in (-1, 1):
            cx, cy = sx * spacing_x / 2, sy * spacing_y / 2
            xmin, xmax = (-width / 2, 0) if sx < 0 else (0, width / 2)
            ymin, ymax = (-depth / 2, 0) if sy < 0 else (0, depth / 2)
            angles = [2 * math.pi * i / segments for i in range(segments)]
            angles.extend(math.atan2(y - cy, x - cx) % (2 * math.pi) for x in (xmin, xmax) for y in (ymin, ymax))
            angles = sorted(set(round(angle, 12) for angle in angles))
            outer, inner = [], []
            for angle in angles:
                dx, dy = math.cos(angle), math.sin(angle)
                tx = ((xmax if dx > 0 else xmin) - cx) / dx if abs(dx) > 1e-10 else math.inf
                ty = ((ymax if dy > 0 else ymin) - cy) / dy if abs(dy) > 1e-10 else math.inf
                distance = min(tx, ty)
                outer.append((round(cx + distance * dx, 8), round(cy + distance * dy, 8)))
                inner.append((round(cx + radius * dx, 8), round(cy + radius * dy, 8)))
            for i in range(len(angles)):
                j = (i + 1) % len(angles)
                planar.extend([(outer[i], outer[j], inner[j]), (outer[i], inner[j], inner[i])])
    edges: Counter = Counter()
    oriented = {}
    triangles = []
    for a, b, c in planar:
        triangles.extend([((a[0], a[1], thickness), (b[0], b[1], thickness), (c[0], c[1], thickness)),
                          ((c[0], c[1], 0.0), (b[0], b[1], 0.0), (a[0], a[1], 0.0))])
        for p, q in ((a, b), (b, c), (c, a)):
            key = tuple(sorted((p, q)))
            edges[key] += 1
            oriented[key] = (p, q)
    for key, count in edges.items():
        if count == 1:
            a, b = oriented[key]
            a0, a1 = (*a, 0.0), (*a, thickness)
            b0, b1 = (*b, 0.0), (*b, thickness)
            triangles.extend([(a0, b0, b1), (a0, b1, a1)])
        elif count != 2:
            raise RuntimeError("Motor bracket triangulation has a non-manifold planar edge")
    return triangles


def _normal(a: tuple, b: tuple, c: tuple) -> tuple:
    u, v = [b[i] - a[i] for i in range(3)], [c[i] - a[i] for i in range(3)]
    normal = (u[1] * v[2] - u[2] * v[1], u[2] * v[0] - u[0] * v[2], u[0] * v[1] - u[1] * v[0])
    length = math.sqrt(sum(value * value for value in normal))
    if length < 1e-12:
        raise RuntimeError("Motor bracket triangulation produced a degenerate triangle")
    return tuple(value / length for value in normal)


def motor_bracket(payload: dict, workspace: Path) -> dict:
    """Generate a flat motor mounting bracket with four through holes, in mm."""
    width = _number(payload.get("width_mm", 60), "width_mm", 5, 1000)
    depth = _number(payload.get("depth_mm", 60), "depth_mm", 5, 1000)
    thickness = _number(payload.get("thickness_mm", 5), "thickness_mm", 0.5, 100)
    spacing_x = _number(payload.get("screw_spacing_x_mm", 31), "screw_spacing_x_mm", 1, 990)
    spacing_y = _number(payload.get("screw_spacing_y_mm", 31), "screw_spacing_y_mm", 1, 990)
    diameter = _number(payload.get("hole_diameter_mm", 3.5), "hole_diameter_mm", 0.5, 200)
    segments = _integer(payload.get("segments", 96), "segments", 32, 256)
    if segments % 4:
        raise ValueError("segments must be a multiple of four")
    if min(width - spacing_x - diameter, depth - spacing_y - diameter, spacing_x - diameter, spacing_y - diameter) < 1:
        raise ValueError("Screw holes must have at least 0.5 mm edge clearance and at least 1 mm separation")
    triangles = _plate_triangles(width, depth, thickness, spacing_x, spacing_y, diameter, segments)
    directory = _folder(workspace, "motor_bracket")
    try:
        with (directory / "motor_bracket.stl").open("wb") as stream:
            stream.write(b"Omega flat motor mounting bracket; units mm".ljust(80, b"\0"))
            stream.write(struct.pack("<I", len(triangles)))
            for a, b, c in triangles:
                stream.write(struct.pack("<12fH", *_normal(a, b, c), *a, *b, *c, 0))
        scad = f"""// Flat mounting bracket, all dimensions in millimeters. Four through holes.
width = {width}; depth = {depth}; thickness = {thickness};
screw_spacing_x = {spacing_x}; screw_spacing_y = {spacing_y}; hole_diameter = {diameter};
$fn = {segments};
difference() {{
  translate([-width/2, -depth/2, 0]) cube([width, depth, thickness]);
  for (x = [-screw_spacing_x/2, screw_spacing_x/2])
    for (y = [-screw_spacing_y/2, screw_spacing_y/2])
      translate([x, y, -0.1]) cylinder(h=thickness+0.2, d=hole_diameter);
}}
"""
        (directory / "motor_bracket.scad").write_text(scad, encoding="utf-8")
        margin = max(width, depth) * 0.3 + 10
        font = max(width, depth) / 24
        circles = "".join(f'<circle cx="{width / 2 + x}" cy="{depth / 2 + y}" r="{diameter / 2}"/>' for x in (-spacing_x / 2, spacing_x / 2) for y in (-spacing_y / 2, spacing_y / 2))
        svg = f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="{-margin} {-margin} {width + 2 * margin} {depth + 2 * margin}" width="900" height="900">
<rect x="{-margin}" y="{-margin}" width="{width + 2*margin}" height="{depth + 2*margin}" fill="white"/>
<g fill="none" stroke="#163f56" stroke-width="{font/9}"><rect width="{width}" height="{depth}"/>{circles}</g>
<g stroke="#738496" stroke-width="{font/15}" stroke-dasharray="2 1"><path d="M{width/2} {-font*2} V{depth+font*2} M{-font*2} {depth/2} H{width+font*2}"/></g>
<g stroke="#163f56" stroke-width="{font/12}"><path d="M0 {-font*3} H{width} M0 {-font*4} V{-font*2} M{width} {-font*4} V{-font*2} M{width+font*3} 0 V{depth} M{width+font*2} 0 H{width+font*4} M{width+font*2} {depth} H{width+font*4}"/></g>
<g fill="#163f56" font-family="sans-serif" font-size="{font}"><text x="{width/2}" y="{-font*4}" text-anchor="middle">{width:g} mm</text><text x="{width+font*4}" y="{depth/2}">{depth:g}</text>
<text x="0" y="{depth+font*4}">4 x diameter {diameter:g} mm THROUGH</text><text x="0" y="{depth+font*5.5}">Hole spacing: {spacing_x:g} x {spacing_y:g} mm</text><text x="0" y="{depth+font*7}">Thickness: {thickness:g} mm | top view | units mm</text></g></svg>'''
        (directory / "motor_bracket.svg").write_text(svg, encoding="utf-8")
        data = {"shape": "flat four-hole motor mounting bracket", "units": "mm", "width_mm": width,
                "depth_mm": depth, "thickness_mm": thickness, "screw_spacing_x_mm": spacing_x,
                "screw_spacing_y_mm": spacing_y, "hole_diameter_mm": diameter, "triangles": len(triangles),
                "maximum_radial_tessellation_error_mm": diameter / 2 * (1 - math.cos(math.pi / segments)),
                "engineering_validation": "Geometry only; no shaft opening, load analysis, fit compensation, or print-process validation."}
        _json_write(directory / "dimensions.json", data)
    except Exception:
        shutil.rmtree(directory)
        raise
    return _result("Created a four-hole flat motor bracket as SCAD, STL, and dimensioned SVG.", data, directory, workspace)


def _cues(value: Any, maximum: float = 86400) -> list[dict]:
    if not isinstance(value, list) or not 1 <= len(value) <= 2000:
        raise ValueError("cues must contain 1 to 2000 {start, end, text} objects")
    result, previous = [], 0.0
    for cue in value:
        if not isinstance(cue, dict):
            raise ValueError("Every cue must be an object")
        start = _number(cue.get("start"), "cue start", 0, maximum)
        end = _number(cue.get("end"), "cue end", 0, maximum)
        if end - start < 0.001 or start < previous:
            raise ValueError("Cues must be ordered, non-overlapping, and at least one millisecond long")
        text = _text(cue.get("text"), "cue text", 2000)
        if "-->" in text or "\n\n" in text:
            raise ValueError("Cue text cannot contain timestamp arrows or blank lines")
        result.append({"start": start, "end": end, "text": text})
        previous = end
    return result


def _timestamp(seconds: float, decimal: str) -> str:
    milliseconds = int(round(seconds * 1000))
    hours, milliseconds = divmod(milliseconds, 3_600_000)
    minutes, milliseconds = divmod(milliseconds, 60_000)
    seconds, milliseconds = divmod(milliseconds, 1000)
    return f"{hours:02}:{minutes:02}:{seconds:02}{decimal}{milliseconds:03}"


def _write_subtitles(cues: list[dict], directory: Path) -> None:
    for suffix, decimal in (("srt", ","), ("vtt", ".")):
        pieces = ["WEBVTT\n\n"] if suffix == "vtt" else []
        for i, cue in enumerate(cues, 1):
            # Do not interpret user text as subtitle formatting tags.
            text = html.escape(cue["text"], quote=False)
            pieces.append(f"{i}\n{_timestamp(cue['start'], decimal)} --> {_timestamp(cue['end'], decimal)}\n{text}\n\n")
        (directory / f"captions.{suffix}").write_text("".join(pieces), encoding="utf-8")


def subtitles(payload: dict, workspace: Path) -> dict:
    """Create SRT and WebVTT from user-supplied timings; no transcription is implied."""
    cues = _cues(payload.get("cues"))
    directory = _folder(workspace, "subtitles")
    _write_subtitles(cues, directory)
    return _result(f"Created SRT and WebVTT for {len(cues)} supplied cues.", {"cue_count": len(cues), "duration_seconds": cues[-1]["end"]}, directory, workspace)


def _ffmpeg() -> str:
    executable = shutil.which("ffmpeg")
    if executable:
        return executable
    try:
        import imageio_ffmpeg
        executable = imageio_ffmpeg.get_ffmpeg_exe()
        if executable and Path(executable).is_file():
            return executable
    except (ImportError, RuntimeError, OSError):
        pass
    raise ArtifactDependencyError("media.clip requires FFmpeg with libx264 (and libass for subtitle burn-in), or the imageio-ffmpeg deployment package")


def media_clip(payload: dict, workspace: Path) -> dict:
    """Trim a staged local video; optional 9:16 crop and supplied-cue burn-in."""
    input_name = _text(payload.get("input_path"), "input_path", 500)
    relative = Path(input_name)
    if relative.is_absolute() or relative.drive or ".." in relative.parts or ":" in input_name or "\\" in input_name:
        raise ValueError("input_path must be a relative workspace path using forward slashes")
    workspace = Path(workspace).resolve()
    current = workspace
    for part in relative.parts:
        current = current / part
        if current.is_symlink():
            raise ValueError("input_path cannot traverse symlinks")
    source = (workspace / relative).resolve()
    if not source.is_relative_to(workspace) or not source.is_file():
        raise ValueError("input_path must name an existing file inside the workspace")
    if source.suffix.lower() not in {".mp4", ".m4v", ".mkv", ".mov", ".webm", ".avi"}:
        raise ValueError("Supported video files: MP4, M4V, MKV, MOV, WebM, AVI")
    if not 1 <= source.stat().st_size <= 256 * 1024 * 1024:
        raise ValueError("Input video must be nonempty and no larger than 256 MiB")
    start = _number(payload.get("start_seconds", 0), "start_seconds", 0, 86400)
    duration = _number(payload.get("duration_seconds", 30), "duration_seconds", 0.1, 180)
    timeout = _number(payload.get("timeout_seconds", 120), "timeout_seconds", 1, 180)
    vertical = _boolean(payload.get("vertical", False), "vertical")
    cues = _cues(payload["cues"], duration) if "cues" in payload else []
    executable = _ffmpeg()
    directory = _folder(workspace, "clip")
    output = directory / "clip.mp4"
    filters = ["scale=720:1280:force_original_aspect_ratio=increase,crop=720:1280,setsar=1"] if vertical else ["scale=w='min(1920,iw)':h='min(1080,ih)':force_original_aspect_ratio=decrease:force_divisible_by=2,setsar=1"]
    if cues:
        _write_subtitles(cues, directory)
        filters.append("subtitles=filename=captions.srt:force_style='FontSize=22,Outline=2,MarginV=32'")
    limit = 48 * 1024 * 1024
    command = [executable, "-nostdin", "-hide_banner", "-loglevel", "error", "-y", "-max_alloc", "67108864",
               "-threads", "2", "-filter_threads", "2", "-protocol_whitelist", "file", "-format_whitelist", "mov,matroska,webm,avi",
               "-ss", str(start), "-i", str(source), "-t", str(duration), "-map", "0:v:0", "-map", "0:a:0?",
               "-vf", ",".join(filters), "-c:v", "libx264", "-threads", "2", "-preset", "veryfast", "-crf", "24",
               "-maxrate", "3M", "-bufsize", "6M", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "128k",
               "-map_metadata", "-1", "-map_chapters", "-1", "-movflags", "+faststart", "-fs", str(limit), str(output)]
    try:
        with tempfile.TemporaryFile() as errors:
            try:
                process = subprocess.run(command, cwd=directory, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                                         stderr=errors, timeout=timeout, check=False, shell=False)
            except subprocess.TimeoutExpired as exc:
                raise RuntimeError(f"FFmpeg exceeded the {timeout:g}-second execution deadline; request a shorter clip") from exc
            errors.seek(0, 2)
            errors.seek(max(0, errors.tell() - 6000))
            detail = errors.read().decode("utf-8", errors="replace").strip()
        if process.returncode:
            if "No such filter: 'subtitles'" in detail or "Unknown encoder" in detail:
                raise ArtifactDependencyError("The deployed FFmpeg lacks the requested encoder or subtitle filter. " + detail[-1500:])
            raise RuntimeError("FFmpeg could not process this local video: " + (detail[-1500:] or f"exit {process.returncode}"))
        if not output.exists() or output.stat().st_size < 1024:
            raise ValueError("The requested interval contains no usable video; check start_seconds against the source duration")
        if output.stat().st_size >= limit - 65536:
            raise ValueError("Clip reached the 48 MiB output limit; request a shorter duration")
    except Exception:
        shutil.rmtree(directory)
        raise
    return _result("Created an MP4 clip" + (" with burned-in supplied captions." if cues else "."), {
        "start_seconds": start, "requested_duration_seconds": duration, "vertical": vertical,
        "size_bytes": output.stat().st_size, "cue_count": len(cues),
        "timing_note": "Cue timings are relative to the output clip; output may end earlier when the source ends.",
    }, directory, workspace)


CAPABILITIES: dict[str, Callable[[dict, Path], dict]] = {
    "survey.bundle": survey_bundle,
    "cad.motor_bracket": motor_bracket,
    "media.clip": media_clip,
    "media.subtitles": subtitles,
}

CAPABILITY_INFO = {
    "survey.bundle": {"description": "Local Google Forms API JSON, Apps Script, standalone survey, and XLSX analytics; no remote creation.",
                      "example": {"title": "Employee feedback", "questions": [{"id": "rating", "title": "Overall satisfaction", "type": "scale", "min": 1, "max": 5}], "responses": [{"rating": 4}]}},
    "cad.motor_bracket": {"description": "Flat rectangular motor mounting bracket with four through screw holes, SCAD/STL/SVG in mm.",
                          "example": {"width_mm": 60, "depth_mm": 60, "thickness_mm": 5, "screw_spacing_x_mm": 31, "screw_spacing_y_mm": 31, "hole_diameter_mm": 3.5}},
    "media.clip": {"description": "Trim a staged workspace video with optional 9:16 crop and supplied caption burn-in; FFmpeg required.",
                   "example": {"input_path": "input.mp4", "start_seconds": 0, "duration_seconds": 30, "vertical": True}},
    "media.subtitles": {"description": "Generate SRT and WebVTT from ordered supplied cues (no automatic transcription).",
                        "example": {"cues": [{"start": 0, "end": 2, "text": "Welcome"}]}},
}
